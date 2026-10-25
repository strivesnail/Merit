#!/usr/bin/env python3
"""Render the editable replacement slide to PDF so it can be reviewed."""

import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import Circle, FancyArrowPatch, FancyBboxPatch, PathPatch, Rectangle
from matplotlib.path import Path as MplPath

CJK_FONT = "/home/jianz/.fonts/wqy-microhei.ttc"
if os.path.isfile(CJK_FONT):
    font_manager.fontManager.addfont(CJK_FONT)
    CJK_FAMILY = font_manager.FontProperties(fname=CJK_FONT).get_name()
else:
    CJK_FAMILY = "DejaVu Sans"
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

PPTX = "/home/jianz/workload/code/NeighborCache/docs/figures/neighborcache_cache_replacement.pptx"
PDF = "/home/jianz/workload/code/NeighborCache/docs/figures/neighborcache_cache_replacement.pdf"
PNG = "/home/jianz/workload/code/NeighborCache/docs/figures/neighborcache_cache_replacement.png"


def rgb_of(color):
    if color is None:
        return "#1a1a1a"
    return f"#{color}"


def fill_hex(shape):
    try:
        fill = shape.fill
        if fill.type is None:
            return None
        return rgb_of(fill.fore_color.rgb)
    except Exception:
        return None


def line_hex(shape):
    try:
        if shape.line.fill.type is None:
            return None
        return rgb_of(shape.line.color.rgb)
    except Exception:
        return "#222222"


def line_dashed(shape):
    try:
        return shape.line.dash_style is not None
    except Exception:
        return False


def line_width(shape):
    try:
        pt = shape.line.width.pt
        return max(pt, 0.6)
    except Exception:
        return 1.0


def has_arrow(shape):
    return shape._element.find(".//" + qn("a:tailEnd")) is not None


def connector_ends(shape):
    xfrm = shape._element.find(".//" + qn("a:xfrm"))
    flip_h = xfrm.get("flipH") == "1"
    flip_v = xfrm.get("flipV") == "1"
    left = shape.left.inches
    top = shape.top.inches
    right = left + shape.width.inches
    bottom = top + shape.height.inches
    start = (right if flip_h else left, bottom if flip_v else top)
    end = (left if flip_h else right, top if flip_v else bottom)
    return start, end


def paragraph_text(paragraph):
    return "".join(run.text for run in paragraph.runs)


def paragraph_style(paragraph):
    size = 12
    bold = False
    color = "#1a1a1a"
    for run in paragraph.runs:
        if run.font.size is not None:
            size = run.font.size.pt
        if run.font.bold:
            bold = True
        try:
            if run.font.color is not None and run.font.color.rgb is not None:
                color = rgb_of(run.font.color.rgb)
        except Exception:
            pass
        if run.text:
            break
    align = paragraph.alignment
    ha = "center"
    if align == PP_ALIGN.LEFT:
        ha = "left"
    elif align == PP_ALIGN.RIGHT:
        ha = "right"
    return size, bold, color, ha


def draw_text_block(ax, x, y, w, h, shape, slide_h):
    if not shape.has_text_frame:
        return
    lines = []
    size, bold, color, ha = 12, False, "#1a1a1a", "center"
    for paragraph in shape.text_frame.paragraphs:
        text = paragraph_text(paragraph)
        if text == "" and not lines:
            continue
        lines.append(text)
        size, bold, color, ha = paragraph_style(paragraph)
    text = "\n".join(lines).strip()
    if not text:
        return
    cx = x + w / 2
    cy = slide_h - (y + h / 2)
    if ha == "left":
        cx = x + 0.04
    elif ha == "right":
        cx = x + w - 0.04
    ax.text(
        cx,
        cy,
        text,
        ha=ha,
        va="center",
        fontsize=size,
        color=color,
        fontweight="bold" if bold else "regular",
        fontfamily=CJK_FAMILY if any(ord(ch) >= 0x2E80 for ch in text) else "DejaVu Sans",
        linespacing=1.05,
        zorder=5,
    )


def draw_table(ax, shape, slide_h):
    table = shape.table
    x0 = shape.left.inches
    y0 = shape.top.inches
    width = shape.width.inches
    height = shape.height.inches
    rows = len(table.rows)
    cols = len(table.columns)
    col_w = [table.columns[i].width.inches for i in range(cols)]
    row_h = height / rows
    y = y0
    for r in range(rows):
        x = x0
        for c in range(cols):
            cell = table.cell(r, c)
            ax.add_patch(
                Rectangle(
                    (x, slide_h - (y + row_h)),
                    col_w[c],
                    row_h,
                    facecolor="white",
                    edgecolor="#222222",
                    linewidth=0.7,
                    zorder=3,
                )
            )
            text = cell.text_frame.paragraphs[0]
            label = paragraph_text(text)
            size, bold, color, _ha = paragraph_style(text)
            ax.text(
                x + col_w[c] / 2,
                slide_h - (y + row_h / 2),
                label,
                ha="center",
                va="center",
                fontsize=size,
                color=color,
                fontweight="bold" if bold else "regular",
                fontfamily="DejaVu Sans",
                zorder=5,
            )
            x += col_w[c]
        y += row_h


def render(pptx=PPTX, pdf=PDF, png=PNG):
    prs = Presentation(pptx)
    slide = prs.slides[0]
    slide_w = prs.slide_width.inches
    slide_h = prs.slide_height.inches
    fig, ax = plt.subplots(figsize=(slide_w, slide_h))
    fig.set_layout_engine(None)
    ax.set_position([0, 0, 1, 1])
    ax.set_xlim(0, slide_w)
    ax.set_ylim(0, slide_h)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")

    for shape in slide.shapes:
        x = shape.left.inches
        y = shape.top.inches
        w = shape.width.inches
        h = shape.height.inches
        if shape.shape_type == MSO_SHAPE_TYPE.LINE:
            (x1, y1), (x2, y2) = connector_ends(shape)
            color = line_hex(shape) or "#1a1a1a"
            dashed = line_dashed(shape)
            style = (0, (3.2, 1.6)) if dashed else "solid"
            geom = shape._element.find(".//" + qn("a:prstGeom"))
            if geom is not None and geom.get("prst", "").startswith("curvedConnector"):
                xm = (x1 + x2) / 2
                path = MplPath(
                    [(x1, slide_h - y1), (xm, slide_h - y1), (xm, slide_h - y2), (x2, slide_h - y2)],
                    [MplPath.MOVETO, MplPath.CURVE4, MplPath.CURVE4, MplPath.CURVE4],
                )
                if has_arrow(shape):
                    ax.add_patch(FancyArrowPatch(path=path, arrowstyle="-|>", mutation_scale=11,
                                                 color=color, linewidth=line_width(shape),
                                                 linestyle=style, zorder=2))
                else:
                    ax.add_patch(PathPatch(path, facecolor="none", edgecolor=color,
                                           linewidth=line_width(shape), linestyle=style, zorder=2))
                continue
            ax.add_patch(
                FancyArrowPatch(
                    (x1, slide_h - y1),
                    (x2, slide_h - y2),
                    arrowstyle="-|>" if has_arrow(shape) else "-",
                    mutation_scale=11,
                    linewidth=line_width(shape),
                    color=color,
                    linestyle=style,
                    shrinkA=0,
                    shrinkB=0,
                    zorder=6 if shape.name.startswith("cross") else 2,
                    clip_on=False,
                )
            )
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            draw_table(ax, shape, slide_h)
            continue

        fill = fill_hex(shape)
        edge = line_hex(shape)
        dashed = line_dashed(shape)
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.auto_shape_type == MSO_SHAPE.OVAL:
            ax.add_patch(
                Circle(
                    (x + w / 2, slide_h - (y + h / 2)),
                    radius=w / 2,
                    facecolor=fill or "#cfcfcf",
                    edgecolor=edge or "#222222",
                    linewidth=0.9,
                    zorder=3,
                )
            )
        elif fill is not None or edge is not None:
            radius = 0.08 if (
                shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                and shape.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
            ) else 0
            patch_y = slide_h - (y + h)
            if radius:
                ax.add_patch(
                    FancyBboxPatch(
                        (x, patch_y),
                        w,
                        h,
                        boxstyle=f"round,pad=0,rounding_size={radius}",
                        facecolor=fill or "none",
                        edgecolor=edge or "none",
                        linewidth=0.9,
                        linestyle=(0, (2.4, 1.4)) if dashed else "solid",
                        zorder=3,
                    )
                )
            else:
                ax.add_patch(
                    Rectangle(
                        (x, patch_y),
                        w,
                        h,
                        facecolor=fill if fill is not None else "none",
                        edgecolor=edge if edge is not None else "none",
                        linewidth=0.9 if edge else 0,
                        linestyle=(0, (2.4, 1.4)) if dashed else "solid",
                        zorder=3,
                    )
                )
        draw_text_block(ax, x, y, w, h, shape, slide_h)

    fig.savefig(pdf, dpi=200)
    fig.savefig(png, dpi=160)
    plt.close(fig)
    print(pdf)
    print(png)


if __name__ == "__main__":
    import sys

    if len(sys.argv) == 2:
        stem = os.path.splitext(sys.argv[1])[0]
        render(sys.argv[1], stem + ".pdf", stem + ".png")
    else:
        render()
