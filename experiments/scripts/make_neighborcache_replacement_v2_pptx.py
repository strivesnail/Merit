#!/usr/bin/env python3
"""Editable PowerPoint of the NEIGHBORCACHE cache replacement figure, paper layout.

N-cache and M-cache are both LRU. Pending pages wait in a commit queue that
holds several replacement tasks before the asynchronous write.
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from make_neighborcache_replacement_pptx import (
    BLUE, INK, MUTED, PALE, SLOT, SLOT_ALT, WHITE, EDGE,
    _run, add_box, add_chip, add_line, add_text, rgb, set_cell_border,
)
from lxml import etree
from pptx.oxml.ns import qn

OUT = "/home/jianz/workload/code/NeighborCache/docs/figures/neighborcache_cache_replacement_v2.pptx"

RED = "C00000"
CROSS = "7F7F7F"
ZOOM_FILL = "EAF1F6"
ZOOM_EDGE_FILL = "BFD3E3"
S = 0.40
GAP = 0.05


def chip_row(slide, x, y, labels, prefix, faces=None):
    for i, label in enumerate(labels):
        face = faces[i] if faces else (SLOT_ALT if i % 2 == 0 else SLOT)
        add_chip(slide, x + i * (S + GAP), y, label, seed=(i == 0 and label != ""),
                 face=face, size=S, name=f"{prefix} {i}")
    return x + len(labels) * (S + GAP) - GAP


def circle(slide, cx, cy, label, r=0.17):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r)
    )
    shape.name = f"heap {label}"
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(SLOT_ALT)
    shape.line.color.rgb = rgb(EDGE)
    tf = shape.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    tf._txBody.find(qn("a:bodyPr")).set("anchor", "ctr")
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, label, 10, INK, True)


def curve(slide, x1, y1, x2, y2, name):
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR

    shape = slide.shapes.add_connector(MSO_CONNECTOR.CURVE, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.name = name
    shape.line.color.rgb = rgb(MUTED)
    shape.line.width = Pt(1.0)
    shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shape


def cross(slide, x, y, w, h, name):
    add_line(slide, x, y, x + w, y + h, color=CROSS, width=1.5, name=f"cross {name} 1")
    add_line(slide, x, y + h, x + w, y, color=CROSS, width=1.5, name=f"cross {name} 2")


def mcache_table(slide, x, y, w, row_h):
    rows = [
        ("ID", "H", "W"),
        ("A", "14", "z:7,  v:4,  w:3"),
        ("B", "5", "r:3,  s:2"),
        ("…", "…", "…"),
        ("q", "1", "y:1"),
    ]
    frame = slide.shapes.add_table(len(rows), 3, Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    frame.name = "M-cache table"
    table = frame.table
    table.columns[0].width = Inches(w * 0.18)
    table.columns[1].width = Inches(w * 0.16)
    table.columns[2].width = Inches(w * 0.66)
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag == qn("a:tableStyleId"):
            tblPr.remove(child)
    for r, rec in enumerate(rows):
        table.rows[r].height = Inches(row_h)
        for c, value in enumerate(rec):
            cell = table.cell(r, c)
            cell.text = ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
                setattr(cell, side, Inches(0.02))
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            color = RED if (r == 1 and c == 1) else INK
            _run(p, value, 12 if r == 0 else 11, color, bold=(r == 0))
            set_cell_border(cell, "222222", 0.75)
    return frame


def build():
    prs = Presentation()
    prs.slide_width = Inches(12.4)
    prs.slide_height = Inches(6.5)
    prs.core_properties.title = "NEIGHBORCACHE cache replacement and update"
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ---------- N-cache (LRU) ----------
    ncx = 1.30
    add_text(slide, ncx + S / 2 - 1.0, 0.00, 2.0, 0.34, "N-cache", 18, True, name="N-cache title")
    add_text(slide, ncx + S / 2 - 1.0, 0.36, 2.0, 0.24, "LRU", 12, False, MUTED, name="N-cache LRU")
    nc_y = {}
    for i, (y, lab) in enumerate(zip([0.65, 1.10, 1.55, 2.00, 2.45], "zvwAu")):
        add_chip(slide, ncx, y, lab, seed=(lab == "A"), face=SLOT_ALT if i % 2 else SLOT,
                 size=S, name=f"N-cache {lab}")
        nc_y[lab] = y
    add_text(slide, ncx, 2.88, S, 0.30, "⋮", 14, True, name="N-cache dots")
    add_chip(slide, ncx, 3.23, "p", face=SLOT_ALT, size=S, name="N-cache p")
    cross(slide, ncx, 3.23, S, S, "p")

    add_line(slide, 0.55, 0.85, ncx - 0.02, 0.85, arrow=True, name="z into N-cache head")
    add_text(slide, 0.62, 0.56, 0.40, 0.24, "(1)", 11, True, BLUE, name="step 1")

    # ---------- M-cache (LRU) ----------
    tx, tw = 8.0, 3.6
    add_text(slide, tx, 0.00, tw, 0.34, "M-cache", 18, True, name="M-cache title")
    add_text(slide, tx, 0.36, tw, 0.24, "LRU", 12, False, MUTED, name="M-cache LRU")
    row_h = 0.30
    ty = 0.62
    mcache_table(slide, tx, ty, tw, row_h)
    strike_y = ty + 4.5 * row_h
    add_line(slide, tx + 0.05, strike_y, tx + tw - 0.05, strike_y, color=CROSS, width=1.5,
             name="cross q row")

    add_chip(slide, 6.60, ty + 1.5 * row_h - S / 2, "z", face=WHITE, size=S, name="access z (2)")
    add_line(slide, 7.00, ty + 1.5 * row_h, tx - 0.02, ty + 1.5 * row_h, arrow=True, name="z into M-cache")
    add_text(slide, 7.25, ty + 1.5 * row_h - 0.28, 0.45, 0.24, "(2)", 11, True, BLUE, name="step 2")

    # ---------- heaps ----------
    add_text(slide, 8.10, 2.25, 1.30, 0.24, "Max Heap", 11, True, name="Max label")
    circle(slide, 8.75, 2.75, "A")
    circle(slide, 8.40, 3.28, "D")
    circle(slide, 9.10, 3.28, "E")
    add_line(slide, 8.65, 2.90, 8.47, 3.13, name="heap A-D")
    add_line(slide, 8.85, 2.90, 9.03, 3.13, name="heap A-E")
    add_box(slide, 8.35, 3.60, 0.80, 0.28, "NonSeed", size=10, bold=False, rounded=True,
            name="NonSeed tag")

    add_text(slide, 10.30, 2.25, 1.30, 0.24, "Min Heap", 11, True, name="Min label")
    circle(slide, 10.95, 2.75, "B")
    circle(slide, 10.60, 3.28, "C")
    circle(slide, 11.30, 3.28, "F")
    add_line(slide, 10.85, 2.90, 10.67, 3.13, name="heap B-C")
    add_line(slide, 11.05, 2.90, 11.23, 3.13, name="heap B-F")
    add_box(slide, 10.55, 3.60, 0.80, 0.28, "Seed", size=10, bold=False, rounded=True,
            name="Seed tag")

    add_text(slide, 9.35, 2.55, 1.20, 0.24, "H(A) > H(B)", 11, True, BLUE, name="heap condition")
    add_text(slide, 9.70, 2.85, 0.50, 0.24, "(3)", 11, True, BLUE, name="step 3")

    # ---------- commit queue: one small block per pending page ----------
    head_x = 3.40
    page_w = 4 * S + 3 * GAP
    wx = head_x + page_w / 2
    blk, step, n_blk = 0.30, 0.40, 6
    by = 3.20
    first_x = wx - blk / 2
    add_box(slide, first_x - 0.15, by - 0.08, (n_blk - 1) * step + blk + 0.30, blk + 0.16, "",
            rounded=True, fill="F3F6F8", name="commit queue box")
    for i in range(n_blk):
        add_box(slide, first_x + i * step, by, blk, blk, "",
                fill=ZOOM_EDGE_FILL if i == n_blk - 1 else SLOT_ALT, name=f"queue block {i}")
    add_text(slide, 2.30, by + 0.02, 1.50, 0.26, "Commit queue", 11, False, MUTED, align="right",
             name="commit queue label")
    add_text(slide, wx + 0.04, by + blk + 0.08, 0.60, 0.22, "head", 10, False, MUTED,
             align="left", name="queue head label")
    tail_x = first_x + (n_blk - 1) * step
    add_text(slide, tail_x - 0.15, by + blk + 0.10, 0.60, 0.22, "tail", 10, False, MUTED,
             name="queue tail label")

    # Zoom into the tail block: the page A z v w just enqueued.
    zoom_page_x = tail_x + blk / 2 - page_w / 2
    zx0, zx1, zy0, zy1 = zoom_page_x - 0.10, zoom_page_x + page_w + 0.10, 2.15, 2.75
    add_box(slide, zx0, zy0, zx1 - zx0, zy1 - zy0, "", rounded=True, fill=ZOOM_FILL, name="zoom frame")
    chip_row(slide, zoom_page_x, zy0 + 0.10, list("Azvw"), "zoom page")
    add_line(slide, tail_x, by, zx0 + 0.02, zy1, color=MUTED, dashed=True, width=0.9, name="zoom left")
    add_line(slide, tail_x + blk, by, zx1 - 0.02, zy1, color=MUTED, dashed=True, width=0.9, name="zoom right")

    queue_right = tail_x + blk + 0.15
    add_line(slide, 8.56, 2.80, queue_right + 0.02, by + blk / 2, arrow=True, name="enqueue")
    add_text(slide, 7.70, 2.70, 0.40, 0.22, "(4)", 11, True, BLUE, name="step 4")

    mx, my = 3.90, (zy0 + zy1) / 2
    for lab in "zvwA":
        curve(slide, ncx + S + 0.04, nc_y[lab] + S / 2, mx, my, name=f"N-cache {lab} to merge")
    add_line(slide, mx, my, zx0 - 0.02, my, dashed=True, arrow=True, color=MUTED, width=1.0,
             name="merged partners to tail page")
    add_text(slide, 4.20, my - 0.30, 0.40, 0.24, "(5)", 11, True, BLUE, name="step 5")

    # ---------- divider ----------
    add_line(slide, 0.30, 4.05, 12.10, 4.05, color=BLUE, dashed=True, width=1.5, name="divider")

    # ---------- D-cache ----------
    add_text(slide, 4.60, 4.10, 2.40, 0.36, "D-cache", 18, True, name="D-cache title")
    bx, ax_ = 1.00, head_x
    for lab, y in (("Page1", 4.60), ("Page2", 5.15)):
        add_text(slide, 0.15, y, 0.80, S, lab, 12, False, INK, align="left", name=f"{lab} label")
    add_text(slide, 0.15, 5.65, 0.80, 0.30, "Pagen", 12, False, INK, align="left", name="Pagen label")

    br = chip_row(slide, bx, 4.60, list("Brst"), "before page1")
    chip_row(slide, ax_, 4.60, list("Azvw"), "after page1")
    add_line(slide, br + 0.08, 4.80, ax_ - 0.08, 4.80, arrow=True, name="page1 before to after")
    chip_row(slide, bx, 5.15, list("Cjkl"), "before page2")
    chip_row(slide, ax_, 5.15, list("Cjkl"), "after page2")
    add_line(slide, br + 0.08, 5.35, ax_ - 0.08, 5.35, arrow=True, name="page2 before to after")
    add_text(slide, bx, 5.62, page_w, 0.30, "…", 14, True, name="before dots")
    add_text(slide, ax_, 5.62, page_w, 0.30, "…", 14, True, name="after dots")
    add_text(slide, bx, 6.08, page_w, 0.28, "Before", 12, False, MUTED, name="Before label")
    add_text(slide, ax_, 6.08, page_w, 0.28, "After", 12, False, MUTED, name="After label")

    add_line(slide, wx, by + blk, wx, 4.58, arrow=True, name="async write")
    add_text(slide, wx - 0.85, 3.58, 0.78, 0.42, "Async\nwrite", 11, False, INK, align="right",
             name="async write label")
    add_text(slide, wx + 0.08, 3.80, 0.40, 0.22, "(6)", 11, True, BLUE, name="step 6")

    ux = 6.60
    ur = chip_row(slide, ux, 4.60, list("Azvw"), "update before")
    add_text(slide, ur + 0.03, 4.60, 0.22, S, "+", 14, True, name="plus")
    add_chip(slide, ur + 0.28, 4.60, "x", face=SLOT_ALT, size=S, name="update x")
    add_line(slide, ur + 0.76, 4.80, ur + 1.10, 4.80, arrow=True, name="update arrow")
    ur2 = chip_row(slide, ur + 1.18, 4.60, list("Azvwx"), "update after")
    add_text(slide, ux, 5.20, ur2 - ux, 0.30, "Update page on insertion", 12, False, INK,
             name="update label")

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "NEIGHBORCACHE 的缓存替换与更新。一次访问碰到 A 的邻居 z，z 进入 N-cache 的 LRU 结构头部，"
        "并记入 M-cache，于是 A 的出边热度 H(A) 和 z 的权重一起升高。如果 N-cache 已满，就按 LRU 去掉最后一位 p；"
        "M-cache 满时同样按 LRU 淘汰非种子条目 q。A 是最热的非种子，B 是最冷的种子。"
        "如果 D-cache 已满，并且 H(A) > H(B)，就用 A 换下 B。新页由 A 和它在 N-cache 中权重最高的邻居 z、v、w 组成，"
        "与其他待写页一起进入提交队列，异步写入后替换掉 B 那一页。如果之后又访问到 A 的邻居 x，就把 x 追加到 A 所在的 page。"
    )

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
