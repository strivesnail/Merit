#!/usr/bin/env python3
"""Editable PowerPoint of the NEIGHBORCACHE cache replacement figure.

Every node, label, and arrow is a native shape. Nothing is a pasted picture.
"""

import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = "/home/jianz/workload/code/NeighborCache/docs/figures/neighborcache_cache_replacement.pptx"

INK = "1A1A1A"
MUTED = "3F3F3F"
BLUE = "1F5F8B"
SEED = "2A2A2A"
SLOT = "CFCFCF"
SLOT_ALT = "E4E4E4"
PALE = "F7F7F7"
WHITE = "FFFFFF"
EDGE = "222222"
LRU_FILL = "F3F6F8"


def rgb(value):
    return RGBColor.from_string(value)


def _anchor(shape):
    body = shape.text_frame._txBody.find(qn("a:bodyPr"))
    body.set("anchor", "ctr")


def _run(paragraph, text, size, color, bold=False, font="Calibri"):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = etree.SubElement(rpr, qn(tag))
        node.set("typeface", font)
    return run


def add_text(slide, x, y, w, h, text, size=12, bold=False, color=INK, align="center", name="text", font="Calibri"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    _anchor(box)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}[align]
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        _run(p, line, size, color, bold, font)
    return box


def add_box(slide, x, y, w, h, text="", size=14, bold=True, fill=WHITE, font=INK, dashed=False,
            rounded=False, name="box"):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(EDGE)
    shape.line.width = Pt(1.0)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if rounded:
        shape.adjustments[0] = 0.18
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    _anchor(shape)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    if text:
        _run(p, text, size, font, bold)
    return shape


def add_chip(slide, x, y, label, seed=False, face=None, dashed=False, size=0.44, name=None):
    fill = SEED if seed else (SLOT if face is None else face)
    font = WHITE if seed else INK
    return add_box(
        slide, x, y, size, size, label, size=15, bold=True, fill=fill, font=font,
        dashed=dashed, name=name or f"chip {label}",
    )


def add_circle(slide, cx, cy, label, r=0.18, name=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r)
    )
    shape.name = name or f"heap {label}"
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(SLOT)
    shape.line.color.rgb = rgb(EDGE)
    shape.line.width = Pt(1.0)
    tf = shape.text_frame
    tf.word_wrap = False
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    _anchor(shape)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, label, 11, INK, True)
    return shape


def add_line(slide, x1, y1, x2, y2, color=INK, dashed=False, arrow=False, width=1.25, name="line"):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.name = name
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = shape.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return shape


def chip_row(slide, x, y, labels, size=0.44, gap=0.06, prefix="chip"):
    for i, label in enumerate(labels):
        seed = i == 0
        face = SLOT if i % 2 == 0 else SLOT_ALT
        add_chip(slide, x + i * (size + gap), y, label, seed=seed, face=face, size=size, name=f"{prefix} {label}")
    return x + len(labels) * (size + gap) - gap


def empty_row(slide, x, y, n, size=0.44, gap=0.06, prefix="empty"):
    for i in range(n):
        add_chip(slide, x + i * (size + gap), y, "", seed=False, face=WHITE, size=size, name=f"{prefix} {i + 1}")
    return x + n * (size + gap) - gap


def set_cell_border(cell, color="222222", width_pt=0.75):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        for old in tcPr.findall(qn(f"a:{edge}")):
            tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(f"a:{edge}"))
        ln.set("w", str(int(width_pt * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        srgb = etree.SubElement(fill, qn("a:srgbClr"))
        srgb.set("val", color)


def add_table(slide, x, y, w, h):
    frame = slide.shapes.add_table(4, 3, Inches(x), Inches(y), Inches(w), Inches(h))
    frame.name = "M-cache table"
    table = frame.table
    table.columns[0].width = Inches(w * 0.18)
    table.columns[1].width = Inches(w * 0.16)
    table.columns[2].width = Inches(w * 0.66)
    # Drop the blue theme so the cells stay editable and plain.
    tbl = table._tbl
    tblPr = tbl.tblPr
    for child in list(tblPr):
        if child.tag == qn("a:tableStyleId"):
            tblPr.remove(child)
    rows = [
        ("ID", "H", "W"),
        ("A", "9", "u:6,  v:4,  w:3"),
        ("B", "2", "r:3,  s:2"),
        ("…", "…", "…"),
    ]
    for r, rec in enumerate(rows):
        for c, value in enumerate(rec):
            cell = table.cell(r, c)
            cell.text = ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            _run(p, value, 12 if r == 0 else 11, INK, bold=(r == 0))
            set_cell_border(cell, "222222", 1.0 if r == 0 else 0.75)
    return frame


def build():
    prs = Presentation()
    prs.slide_width = Inches(16.4)
    prs.slide_height = Inches(5.65)
    prs.core_properties.title = "NEIGHBORCACHE cache replacement and update"
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    s = 0.44
    gap = 0.06

    # ---------- N-cache ----------
    add_text(slide, 0.55, 0.02, 2.4, 0.34, "N-cache", 20, True, name="N-cache title")
    nx = [0.30, 0.80, 1.30, 1.80, 2.70]
    ny = 0.95
    for i, (x, lab) in enumerate(zip(nx, "Auvwp")):
        add_chip(
            slide, x, ny, lab, seed=(lab == "A"),
            face=SLOT if i % 2 == 0 else SLOT_ALT, size=s,
            name=f"N-cache {lab}",
        )
    add_text(slide, 0.18, 0.66, 0.68, 0.26, "MRU", 12, True, BLUE, name="MRU label")
    add_text(slide, 2.58, 0.66, 0.68, 0.26, "LRU", 12, True, BLUE, name="LRU label")

    add_chip(slide, 0.30, 1.78, "y", face=PALE, size=0.40, name="admit y")
    add_line(slide, 0.50, 1.78, 0.50, 1.39, arrow=True, name="admit arrow")
    add_text(slide, 0.82, 1.82, 0.42, 0.28, "(1)", 13, True, BLUE, name="step 1")

    # ---------- M-cache ----------
    add_text(slide, 6.45, 0.10, 3.40, 0.36, "M-cache", 20, True, name="M-cache title")
    # One access is recorded into M-cache. It does not leave the N-cache row.
    add_chip(slide, 5.05, 0.22, "z", face=PALE, size=0.40, name="access z")
    add_text(slide, 5.50, 0.16, 0.46, 0.26, "(2)", 13, True, BLUE, name="step 2")
    add_line(slide, 5.45, 0.48, 6.42, 0.95, color=BLUE, dashed=True, arrow=True, name="access to mcache")
    add_table(slide, 6.45, 0.52, 3.40, 1.28)
    add_box(slide, 6.85, 1.92, 2.60, 0.40, "", rounded=True, fill=LRU_FILL, name="NonSeed LRU box")
    add_text(slide, 6.90, 1.96, 2.50, 0.32, "NonSeed LRU", 12, False, INK, name="NonSeed LRU")

    # H goes down beside the LRU box, not through it.
    add_line(slide, 6.55, 1.80, 6.55, 3.48, color=BLUE, dashed=True, arrow=True, name="H to max")
    add_text(slide, 6.15, 2.55, 0.36, 0.24, "H", 12, True, BLUE, name="H label")
    add_line(slide, 9.85, 1.16, 9.98, 1.16, color=BLUE, dashed=True, name="H to min jog")
    add_line(slide, 9.98, 1.16, 9.98, 3.48, color=BLUE, dashed=True, name="H to min down")
    add_line(slide, 9.98, 3.48, 9.15, 3.48, color=BLUE, dashed=True, arrow=True, name="H to min")

    # ---------- heaps ----------
    add_text(slide, 5.55, 3.22, 0.55, 0.26, "Max", 13, True, align="right", name="Max label")
    add_chip(slide, 6.15, 3.48, "A", seed=True, size=0.46, name="max root A")
    add_text(slide, 9.15, 2.98, 0.55, 0.24, "Min", 13, True, align="left", name="Min label")
    add_chip(slide, 8.55, 3.48, "B", seed=True, size=0.46, name="min root B")
    add_circle(slide, 6.05, 4.42, "D", name="heap D")
    add_circle(slide, 6.78, 4.42, "E", name="heap E")
    add_circle(slide, 8.45, 4.42, "C", name="heap C")
    add_circle(slide, 9.18, 4.42, "F", name="heap F")
    add_line(slide, 6.38, 3.94, 6.05, 4.24, name="heap A-D")
    add_line(slide, 6.38, 3.94, 6.78, 4.24, name="heap A-E")
    add_line(slide, 8.78, 3.94, 8.45, 4.24, name="heap B-C")
    add_line(slide, 8.78, 3.94, 9.18, 4.24, name="heap B-F")

    add_text(slide, 6.85, 3.12, 1.55, 0.24, "(3)", 13, True, BLUE, name="step 3")
    add_text(slide, 6.62, 3.58, 1.85, 0.26, "H(A) > H(B)", 12, True, BLUE, name="heap condition")

    # ---------- D-cache ----------
    add_text(slide, 11.3, 0.12, 3.4, 0.36, "D-cache", 20, True, name="D-cache title")
    add_text(slide, 10.15, 0.52, 1.6, 0.24, "Before", 12, False, MUTED, name="Before label")
    add_text(slide, 12.90, 0.52, 1.6, 0.24, "After", 12, False, MUTED, name="After label")
    before_right = chip_row(slide, 10.20, 0.80, list("Brst"), size=s, gap=gap, prefix="before")
    after_x = 12.90
    after_right = chip_row(slide, after_x, 0.80, list("Auvw"), size=s, gap=gap, prefix="after")
    add_line(slide, before_right + 0.10, 1.02, after_x - 0.08, 1.02, arrow=True, name="before to after")
    # x arrives later and is written onto the tail of the page already committed.
    add_chip(slide, 15.00, 1.38, "x", face=PALE, size=0.40, name="later x")
    add_text(slide, 13.10, 1.32, 0.46, 0.26, "(6)", 13, True, BLUE, name="step 6")
    add_text(slide, 13.55, 1.28, 1.40, 0.52, "后续插入\nA 的邻居", 12, False, INK, align="left", name="later insert", font="Microsoft YaHei")
    add_line(slide, 15.00, 1.52, 14.90, 1.18, color=BLUE, arrow=True, name="x appended later")
    chip_row(slide, 10.20, 2.12, list("Cjkl"), size=s, gap=gap, prefix="before C")
    chip_row(slide, after_x, 2.12, list("Cjkl"), size=s, gap=gap, prefix="after C")
    add_line(slide, before_right + 0.10, 2.34, after_x - 0.08, 2.34, color=MUTED, dashed=True, arrow=True, name="C unchanged")

    add_text(slide, 10.20, 2.72, 0.90, 0.24, "Pending", 12, False, MUTED, align="left", name="Pending label")
    pend_x = 11.40
    pend_right = chip_row(slide, pend_x, 2.68, list("Auvw"), size=s, gap=gap, prefix="pending")
    add_line(slide, pend_x + 1.0, 3.16, pend_x + 1.0, 3.40, arrow=True, name="pending to queue")
    empty_row(slide, pend_x, 3.40, 4, size=s, gap=gap, prefix="commit")
    add_text(slide, pend_x, 3.88, 2.10, 0.22, "Commit queue", 11, False, MUTED, name="Commit queue label")

    add_line(slide, 9.55, 3.72, pend_x - 0.08, 3.62, arrow=True, name="heap to commit")
    add_text(slide, 10.05, 3.48, 0.50, 0.24, "(4)", 13, True, BLUE, name="step 4")

    add_line(slide, pend_right + 0.06, 3.62, 15.70, 3.62, name="async horizontal")
    add_line(slide, 15.70, 3.62, 15.70, 1.02, name="async vertical")
    add_line(slide, 15.70, 1.02, after_right + 0.02, 1.02, arrow=True, name="async into after")
    add_text(slide, 15.78, 2.20, 0.58, 0.48, "Async\nwrite", 11, False, INK, name="async write label")
    add_text(slide, 14.95, 0.50, 0.50, 0.24, "(5)", 13, True, BLUE, name="step 5")

    add_text(slide, 3.5, 5.15, 9.4, 0.36, "NEIGHBORCACHE cache replacement and update", 18, True, name="figure title")

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "图上只保留结构和编号，机制见配套说明。\n"
        "(1) N-cache 按新旧排成一行，左端 MRU，右端 LRU。y 从 MRU 端进入，p 在 LRU 端，下一次从这里换出，不进入新页。\n"
        "(2) 一次访问单独记入 M-cache，更新 H 和 W，不是从 N-cache 搬过去。NonSeed LRU 只在元数据满时淘汰条目。\n"
        "最大堆是仍在 N-cache 中的最热非种子，最小堆是最冷种子。\n"
        "(3) D-cache 已满且 H(A) > H(B) 时，用 A 替换 B。圆圈只表示堆序。选中后 A 进入 ReadyToInsertion，B 进入 ReadyToDeletion。\n"
        "Pending 页由 A 和 W(A) 中仍有 payload 的伙伴组成。\n"
        "(4) 该页进入 commit queue，异步写入。\n"
        "(5) 写入提交成功后，A 变为 Seed，B 变为 NonSeed。C 的页不变。\n"
        "(6) x 是 A 的邻居，页已经提交之后才到来，从页外写到 A u v w 的页尾，不是和 u v w 一起生成的。旧成员保留。"
    )

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
