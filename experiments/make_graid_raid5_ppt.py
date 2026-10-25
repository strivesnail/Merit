#!/usr/bin/env python3
"""Native one-slide GRAID RAID5 briefing (English)."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

PPTX_PATH = "/home/jianz/workload/code/NeighborCache/experiments/GRAID_RAID5.pptx"

NAVY = RGBColor(0x10, 0x2A, 0x43)
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
GREEN = RGBColor(0x14, 0x7D, 0x64)
RED = RGBColor(0x9B, 0x2C, 0x2C)
INK = RGBColor(0x10, 0x2A, 0x43)
MUTED = RGBColor(0x48, 0x65, 0x81)
LINE = RGBColor(0xD9, 0xE2, 0xEC)
PAPER = RGBColor(0xF7, 0xF8, 0xFA)
ALT = RGBColor(0xF0, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FOOTER = RGBColor(0x9F, 0xB3, 0xC8)


def rgb_hex(color: RGBColor) -> str:
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def set_run(run, text, size, color, bold=False, name="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_textbox(slide, left, top, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.TOP: "t", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size, color, bold=bold)
    return box


def add_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def add_round_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def set_cell_fill(cell, color: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith("solidFill") or child.tag.endswith("gradFill"):
            tcPr.remove(child)
    solid = etree.SubElement(tcPr, qn("a:solidFill"))
    srgb = etree.SubElement(solid, qn("a:srgbClr"))
    srgb.set("val", rgb_hex(color))


def set_cell_border(cell, color="D9E2EC", width="6350"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        for child in list(tcPr):
            if child.tag.endswith(edge):
                tcPr.remove(child)
        ln = etree.SubElement(tcPr, qn(f"a:{edge}"))
        ln.set("w", width)
        sf = etree.SubElement(ln, qn("a:solidFill"))
        srgb = etree.SubElement(sf, qn("a:srgbClr"))
        srgb.set("val", color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")


def style_cell(cell, text, size, color, bold=False, fill=None, align=PP_ALIGN.CENTER, border="D9E2EC"):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size, color, bold=bold)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.12)
    cell.margin_right = Inches(0.12)
    cell.margin_top = Inches(0.06)
    cell.margin_bottom = Inches(0.06)
    if fill is not None:
        set_cell_fill(cell, fill)
    set_cell_border(cell, color=border)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER

    add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.18), NAVY)
    add_rect(slide, Inches(0), Inches(1.18), Inches(13.333), Inches(0.06), TEAL)
    add_textbox(slide, Inches(0.45), Inches(0.12), Inches(12.4), Inches(0.62), "GRAID RAID5", 32, WHITE, bold=True)
    add_textbox(slide, Inches(0.45), Inches(0.68), Inches(12.4), Inches(0.38), "Setup + fio", 16, FOOTER)

    add_round_rect(slide, Inches(0.45), Inches(1.50), Inches(12.43), Inches(1.72), WHITE, LINE)
    env = [
        ("Host", "Threadripper 7960X  ·  2× RTX PRO 5000  ·  GRAID AE 2.0"),
        ("Storage", "5× 990 PRO 1TB, RAID5"),
        ("Compare", "Same disks, GRAID vs Linux mdadm"),
    ]
    for i, (label, value) in enumerate(env):
        y = Inches(1.58 + i * 0.52)
        add_textbox(slide, Inches(0.70), y, Inches(1.50), Inches(0.48), label, 14, MUTED, bold=True)
        add_textbox(slide, Inches(2.20), y, Inches(10.30), Inches(0.48), value, 18, INK)

    rows = [
        ["", "GRAID", "mdadm", "vs mdadm"],
        ["Seq. write", "23.3 GiB/s", "2.3 GiB/s", "10×"],
        ["Seq. read", "15.2 GiB/s", "27.1 GiB/s", "0.56×"],
        ["Rand. read", "2.1M IOPS", "3.4M IOPS", "0.61×"],
        ["Rand. write", "853k IOPS", "113k IOPS", "7.5×"],
    ]
    table_shape = slide.shapes.add_table(5, 4, Inches(0.45), Inches(3.50), Inches(12.43), Inches(3.40))
    table = table_shape.table
    widths = [Inches(3.40), Inches(3.01), Inches(3.01), Inches(3.01)]
    for i, w in enumerate(widths):
        table.columns[i].width = w

    ratio_color = {1: GREEN, 2: RED, 3: RED, 4: GREEN}
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            if r == 0:
                style_cell(cell=table.cell(r, c), text=text, size=15, color=WHITE, bold=True, fill=NAVY, border="102A43")
            else:
                fill_c = ALT if r % 2 else WHITE
                color = INK
                bold = c == 0
                size = 18
                if c == 3:
                    color = ratio_color[r]
                    bold = True
                style_cell(
                    cell=table.cell(r, c),
                    text=text,
                    size=size,
                    color=color,
                    bold=bold,
                    fill=fill_c,
                    align=PP_ALIGN.CENTER if c else PP_ALIGN.LEFT,
                )

    add_rect(slide, Inches(0), Inches(7.12), Inches(13.333), Inches(0.38), NAVY)
    add_textbox(
        slide,
        Inches(0.45),
        Inches(7.12),
        Inches(12.4),
        Inches(0.38),
        "5× Samsung 990 PRO  ·  RAID5  ·  GRAID vs mdadm",
        12,
        FOOTER,
    )

    prs.save(PPTX_PATH)
    print(PPTX_PATH)


if __name__ == "__main__":
    main()
